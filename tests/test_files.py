import io
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import artifacts as art
from PIL import Image


class ArtifactTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self): self.temp.cleanup()

    def make(self,path,kind,content):
        return art.produce_artifact(self.root,{'path':path,'kind':kind,'content':content})

    def test_invalid_json_leaves_no_file(self):
        with self.assertRaises(ValueError): self.make('bad.json','text','{"broken":')
        self.assertEqual(list(self.root.iterdir()),[])

    def test_invalid_python_is_rejected_without_execution(self):
        with self.assertRaises(SyntaxError): self.make('bad.py','text','def broken(:')
        self.assertEqual(list(self.root.iterdir()),[])

    def test_zip_fails_whole_bundle_on_invalid_entry(self):
        with self.assertRaises(ValueError):
            self.make('bundle.zip','zip',[{'path':'good.txt','kind':'text','content':'hello'}, {'path':'broken.json','kind':'text','content':'{'}])
        self.assertEqual(list(self.root.iterdir()),[])

    def test_zip_traversal_rejected(self):
        with self.assertRaises(ValueError):
            self.make('bundle.zip','zip',[{'path':'../outside.txt','kind':'text','content':'no'}])
        self.assertEqual(list(self.root.iterdir()),[])

    def test_model_is_complete_with_companions(self):
        import cad_engine
        with patch.object(cad_engine,'_render_view',return_value=Image.new('RGB',(64,64))):
            result = self.make('part.stl','stl',{'tree':{'shape':'loft','sections':[[0,0,0,5,5],[0,0,12,4,4]]}})
        expected = {'part.stl','part.glb','part.model.json','part-preview.png','part-validation.json'}
        self.assertEqual({p.name for p in self.root.iterdir()},expected)
        self.assertGreater(result['bytes'],84)
        report = json.loads((self.root/'part-validation.json').read_text())
        self.assertEqual(set(report['exports']),{'part.stl','part.glb'})

    def test_failed_model_leaves_no_companions(self):
        with self.assertRaises(ValueError):
            self.make('bad.stl','stl',{'tree':{'shape':'mesh','vertices':[[0,0,0]],'faces':[]}})
        self.assertEqual(list(self.root.iterdir()),[])

    def test_word_preserves_paragraphs_and_table(self):
        self.make('report.docx','docx','# Test report\n\nA complete paragraph.\n\n| Item | Count |\n|---|---|\n| Widget | 3 |')
        from docx import Document
        doc = Document(self.root/'report.docx')
        self.assertTrue(any('complete paragraph' in p.text for p in doc.paragraphs))
        self.assertEqual(doc.tables[0].cell(1,1).text,'3')

    def test_pdf_is_readable_and_contains_unicode(self):
        self.make('report.pdf','pdf','# Test report\n\nCafé — a complete paragraph.\n\n- First item\n- Second item')
        from pypdf import PdfReader
        text = ''.join(p.extract_text() for p in PdfReader(self.root/'report.pdf').pages)
        self.assertIn('complete paragraph',text)
        self.assertIn('Café',text)

    def test_spreadsheet_preserves_numbers_formulas_and_ids(self):
        self.make('budget.xlsx','xlsx',[['ID','Amount'],['00123',12.5],['Total','=SUM(B2:B2)']])
        from openpyxl import load_workbook
        wb = load_workbook(self.root/'budget.xlsx')
        self.assertEqual(wb.active['A2'].value,'00123')
        self.assertEqual(wb.active['B2'].value,12.5)
        self.assertEqual(wb.active['B3'].value,'=SUM(B2:B2)')
        self.assertTrue(wb.calculation.fullCalcOnLoad)
        wb.close()

    def test_slides_preserve_notes_and_table(self):
        self.make('deck.pptx','pptx',[{'layout':'title','title':'Review','subtitle':'October'}, {'title':'Data','table':[['Item','Count'],['Widgets','3']],'notes':'Discuss the assumptions.'}])
        from pptx import Presentation
        deck = Presentation(self.root/'deck.pptx')
        self.assertEqual(len(deck.slides),2)
        self.assertTrue(any(s.has_table for s in deck.slides[1].shapes))
        self.assertIn('assumptions',deck.slides[1].notes_slide.notes_text_frame.text)

    @unittest.skipUnless(importlib.util.find_spec('matplotlib'), 'matplotlib is not installed in this local test environment')
    def test_chart_is_real_png(self):
        self.make('chart.png','chart',{'type':'bar','title':'Counts','labels':['A','B'],'series':[{'name':'Count','values':[2,4]}]})
        with Image.open(self.root/'chart.png') as im:
            self.assertEqual(im.format,'PNG'); self.assertGreater(im.width,500)

    def test_valid_code_is_saved_without_running(self):
        self.make('safe.py','text','raise RuntimeError("never execute generated code")')
        self.assertIn('RuntimeError',(self.root/'safe.py').read_text())


if __name__ == '__main__': unittest.main()
