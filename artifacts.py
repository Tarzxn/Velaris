"""Velaris file production.

Everything that turns a model's JSON "files" list into real, verified files on
disk lives here: Markdown-lite parsing, DOCX / PDF / PPTX / XLSX renderers,
charts, images, SVG, ZIP, STL, and a validation pass that re-opens each result
with the matching library so a corrupt file is caught before the user sees it.
"""
import base64
import csv
import datetime as dt
import io
import json
import math
import re
import shutil
import tempfile
import textwrap
import uuid
import zipfile
from pathlib import Path
from urllib.parse import quote
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape as xml_escape

import requests

from stl_engine import render_ascii_stl, run_stl_program, safe_path

ARTIFACT_KINDS = {"text", "base64", "image", "chart", "docx", "xlsx", "pptx", "pdf", "stl", "zip"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".svg"}
RASTER_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
OFFICE_KINDS = {"pdf", "docx", "xlsx", "pptx"}

EXT_TO_KIND = {
    ".docx": "docx", ".xlsx": "xlsx", ".pptx": "pptx", ".pdf": "pdf", ".stl": "stl", ".zip": "zip",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".webp": "image", ".gif": "image",
}
KIND_ALIASES = {
    "markdown": "text", "md": "text", "code": "text", "html": "text", "svg": "text", "json": "text",
    "csv": "text", "txt": "text", "plain": "text", "string": "text", "file": "text",
    "word": "docx", "doc": "docx", "document": "docx",
    "excel": "xlsx", "spreadsheet": "xlsx", "sheet": "xlsx",
    "powerpoint": "pptx", "ppt": "pptx", "slides": "pptx", "presentation": "pptx",
    "png": "image", "jpg": "image", "jpeg": "image", "picture": "image", "photo": "image",
    "3d": "stl", "model": "stl", "mesh": "stl",
    "graph": "chart", "plot": "chart",
    "binary": "base64", "bin": "base64", "b64": "base64",
    "archive": "zip",
}
DEFAULT_EXT = {"pdf": ".pdf", "docx": ".docx", "xlsx": ".xlsx", "pptx": ".pptx", "stl": ".stl",
               "image": ".png", "chart": ".png", "zip": ".zip", "text": ".md", "base64": ".bin"}

MAX_BASE64_BYTES = 30 * 1024 * 1024
MAX_ZIP_ENTRIES = 100
MAX_SHEETS = 20
MAX_SHEET_ROWS = 50_000

IMAGE_API_URL = "https://image.pollinations.ai/prompt/{prompt}"
IMAGE_TIMEOUT = 90

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def strip_control(text):
    return _CONTROL_CHARS.sub("", str(text))


# --------------------------------------------------------------------------
# Request analysis
# --------------------------------------------------------------------------
_FORMAT_WORDS = [
    ("pdf", r"pdf|portable document"),
    ("docx", r"docx|word (?:document|doc|file)|ms word|microsoft word"),
    ("xlsx", r"xlsx|excel(?: (?:spreadsheet|file|sheet|workbook))?|spreadsheet|workbook"),
    ("pptx", r"pptx|powerpoint|power point|slide deck|slideshow|slides|presentation|pitch deck"),
    ("stl", r"stl|3d[- ]print(?:able|ing)?(?: (?:model|file))?|3d model"),
    ("zip", r"zip(?: archive| file)?"),
    ("image", r"png|jpe?g|webp|picture|image|illustration|photo(?:graph)?"),
]
_CREATE_VERBS = (r"make|create|generate|build|write|produce|export|save|give|prepare|draft|design|"
                 r"turn|convert|compile|put|output|render|download|need|want")
_NOT_IMAGE = re.compile(r"\b(?:docker|container|disk|iso|vm|machine|base|boot|os|system|virtual)\s+images?\b")


def requested_artifact_kinds(prompt):
    """Ordered list of explicitly requested deliverable formats. Conservative:
    a format word only counts next to a creation verb, a preposition like
    'as a'/'in', or as a file extension."""
    text = _NOT_IMAGE.sub(" ", str(prompt).lower())
    found = []
    for kind, words in _FORMAT_WORDS:
        patterns = [
            rf"\.(?:{words})\b",
            rf"\b(?:{_CREATE_VERBS})\b[^.!?\n]{{0,80}}?\b(?:{words})\b",
            rf"\b(?:as|in|into|to)\s+(?:an?\s+|the\s+)?(?:{words})\b",
            rf"\b(?:{words})\s+(?:file|document|version|export|download)\b",
        ]
        if any(re.search(p, text) for p in patterns):
            found.append(kind)
    return found


def requested_artifact_kind(prompt):
    kinds = requested_artifact_kinds(prompt)
    return kinds[0] if kinds else None


def kind_satisfied(kind, created):
    """created: list of dicts with 'kind' and 'path'."""
    for c in created:
        if c["kind"] == kind:
            return True
        if kind == "image" and (c["kind"] == "chart" or c["path"].lower().endswith(".svg")):
            return True
    return False


# --------------------------------------------------------------------------
# Normalization
# --------------------------------------------------------------------------
_PLACEHOLDER_STEMS = {"generation", "generated", "output", "document", "response", "result", "report",
                      "answer", "content", "file", "untitled"}
_BAD_NAME_CHARS = re.compile(r"[^\w.\-+@&,'() ]+", re.UNICODE)


def sanitize_relpath(raw):
    path = safe_path(str(raw).strip())
    parts = []
    for part in path.parts:
        cleaned = _BAD_NAME_CHARS.sub("_", part).strip(" .")
        cleaned = cleaned.lstrip(".")
        if not cleaned:
            raise ValueError("Unsafe output filename")
        if len(cleaned) > 100:
            stem, dot, ext = cleaned.rpartition(".")
            cleaned = (stem[:90] + dot + ext[:8]) if dot else cleaned[:100]
        parts.append(cleaned)
    return Path(*parts)


def default_extension(kind):
    return DEFAULT_EXT.get(kind, ".bin")


def normalize_artifact_item(item, prompt=""):
    if not isinstance(item, dict):
        raise ValueError("Artifact entry is not an object")
    raw_path = str(item.get("path", "")).strip()
    if not raw_path:
        raise ValueError("Artifact has no filename")
    if item.get("content") is None:
        raise ValueError("Artifact has no content (the response was probably cut off)")
    path = sanitize_relpath(raw_path)
    ext = path.suffix.lower()

    kind = str(item.get("kind", "") or "").lower().strip()
    kind = KIND_ALIASES.get(kind, kind)
    if not kind:
        kind = EXT_TO_KIND.get(ext, "text")
    if kind not in ARTIFACT_KINDS:
        raise ValueError(f"Unsupported artifact kind: {kind}")

    requested = requested_artifact_kinds(prompt)
    only_requested = requested[0] if len(requested) == 1 else None
    if (kind == "text" and ext in {".md", ".txt", ".markdown", ""} and only_requested in OFFICE_KINDS | {"stl", "zip"}
            and path.stem.lower() in _PLACEHOLDER_STEMS):
        # Models often name the real deliverable "generation.md"; the server owns the format.
        kind = only_requested
        path = path.with_suffix(default_extension(kind))
    elif kind == "text" and ext in EXT_TO_KIND and EXT_TO_KIND[ext] != "text":
        kind = EXT_TO_KIND[ext]
    elif kind in OFFICE_KINDS | {"stl", "zip"} and ext != default_extension(kind):
        path = path.with_suffix(default_extension(kind))  # keep extension and renderer consistent

    if kind == "image" and path.suffix.lower() not in RASTER_EXTENSIONS:
        path = path.with_suffix(".png")
    if kind == "chart":
        path = path.with_suffix(".png")
    if kind == "text" and not path.suffix:
        path = path.with_suffix(".txt")
    return {"path": path.as_posix(), "kind": kind, "content": item.get("content", "")}


def _unique_rel(root, rel):
    rel_path = Path(rel)
    candidate, n = rel_path, 2
    while (root / candidate).exists():
        candidate = rel_path.with_name(f"{rel_path.stem}-{n}{rel_path.suffix}")
        n += 1
    return candidate.as_posix()


# --------------------------------------------------------------------------
# Markdown-lite parsing (shared by DOCX / PDF / fallbacks)
# --------------------------------------------------------------------------
_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^\s)]+)\)")
_INLINE = re.compile(r"(\*\*[^*\n]+?\*\*|`[^`\n]+`|(?<![\w*])\*[^*\s][^*\n]*?\*(?![\w*]))")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_HR = re.compile(r"^(?:-{3,}|\*{3,}|_{3,})$")
_NUMERIC_CELL = re.compile(r"^[\s$€£¥+\-(]*\d[\d,.\s]*[%)]?\s*$")


def parse_inline(text):
    """-> [(text, flag)] with flag in '', 'b', 'i', 'c'."""
    text = _LINK.sub(lambda m: m.group(1) if m.group(1) == m.group(2) else f"{m.group(1)} ({m.group(2)})", str(text))
    segments, pos = [], 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            segments.append((text[pos:m.start()], ""))
        token = m.group(0)
        if token.startswith("**"):
            segments.append((token[2:-2], "b"))
        elif token.startswith("`"):
            segments.append((token[1:-1], "c"))
        else:
            segments.append((token[1:-1], "i"))
        pos = m.end()
    if pos < len(text):
        segments.append((text[pos:], ""))
    return segments or [("", "")]


def _split_row(line):
    cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [c.replace("\\|", "|").strip() for c in cells]


def parse_markdown_blocks(content):
    """-> list of blocks:
    ('heading', level, text) | ('bullet', text, depth) | ('number', text, n, depth) |
    ('table', rows, aligns) | ('code', lang, text) | ('quote', text) | ('hr',) | ('para', text)"""
    lines = strip_control(content).replace("\r\n", "\n").replace("\r", "\n").replace("\t", "    ").split("\n")
    blocks, i, n = [], 0, len(lines)
    while i < n:
        raw = lines[i]
        stripped = raw.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            lang, code = stripped[3:].strip(), []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            blocks.append(("code", lang, "\n".join(code)))
            continue
        m = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", stripped)
        if m:
            blocks.append(("heading", min(len(m.group(1)), 4), m.group(2))); i += 1
            continue
        if _HR.match(stripped):
            blocks.append(("hr",)); i += 1
            continue
        if "|" in stripped and i + 1 < n and "|" in lines[i + 1] and _TABLE_SEP.match(lines[i + 1].strip()):
            header = _split_row(stripped)
            aligns = []
            for sep in _split_row(lines[i + 1]):
                aligns.append("center" if sep.startswith(":") and sep.endswith(":") else "right" if sep.endswith(":") else "left")
            rows = [header]
            i += 2
            while i < n and lines[i].strip() and "|" in lines[i]:
                rows.append(_split_row(lines[i])); i += 1
            cols = max(len(r) for r in rows)
            rows = [r + [""] * (cols - len(r)) for r in rows]
            aligns = (aligns + ["left"] * cols)[:cols]
            blocks.append(("table", rows, aligns))
            continue
        m = re.match(r"^(\s*)[-*+•]\s+(.*)", raw)
        if m:
            blocks.append(("bullet", m.group(2).strip(), min(len(m.group(1)) // 2, 3))); i += 1
            continue
        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)", raw)
        if m:
            blocks.append(("number", m.group(3).strip(), int(m.group(2)), min(len(m.group(1)) // 2, 3))); i += 1
            continue
        m = re.match(r"^>\s?(.*)", stripped)
        if m:
            blocks.append(("quote", m.group(1))); i += 1
            continue
        blocks.append(("para", stripped)); i += 1
    return blocks


def _title_from_blocks(blocks, default="Document"):
    for b in blocks:
        if b[0] == "heading":
            return re.sub(r"[*`]", "", b[2])[:120]
    return default


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def _docx_runs(paragraph, text, size=None):
    from docx.shared import Pt, RGBColor
    for seg, flag in parse_inline(text):
        if not seg:
            continue
        run = paragraph.add_run(seg)
        if flag == "b":
            run.bold = True
        elif flag == "i":
            run.italic = True
        elif flag == "c":
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
        if size and flag != "c":
            run.font.size = Pt(size)


def _docx_shade(properties, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    properties.append(shd)


def _docx_border(paragraph, side="bottom", color="9CA3AF", size=8):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single"); edge.set(qn("w:sz"), str(size)); edge.set(qn("w:space"), "4"); edge.set(qn("w:color"), color)
    borders.append(edge)
    pPr.append(borders)


def _docx_page_numbers(section):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)


def write_docx(path, content):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    blocks = parse_markdown_blocks(content)
    if not blocks:
        raise ValueError("Document has no content")
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(0.9)
        _docx_page_numbers(section)
    doc.core_properties.title = _title_from_blocks(blocks)
    doc.core_properties.author = "Velaris"

    for block in blocks:
        tag = block[0]
        if tag == "heading":
            heading = doc.add_heading(level=block[1])
            _docx_runs(heading, block[2])
        elif tag == "bullet":
            style = "List Bullet" if block[2] == 0 else f"List Bullet {min(block[2] + 1, 3)}"
            _docx_runs(doc.add_paragraph(style=style), block[1])
        elif tag == "number":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35 + 0.3 * block[3])
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(3)
            p.add_run(f"{block[2]}.\u00a0")
            _docx_runs(p, block[1])
        elif tag == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            _docx_border(p, "left", "9CA3AF", 18)
            for seg, flag in parse_inline(block[1]):
                run = p.add_run(seg); run.italic = True
                run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
                if flag == "b":
                    run.bold = True
        elif tag == "hr":
            p = doc.add_paragraph(); _docx_border(p)
        elif tag == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.space_after = Pt(8)
            _docx_shade(p._p.get_or_add_pPr(), "F3F4F6")
            lines = block[2].split("\n")
            for idx, line in enumerate(lines):
                run = p.add_run(line)
                run.font.name = "Consolas"; run.font.size = Pt(9)
                if idx < len(lines) - 1:
                    run.add_break()
        elif tag == "table":
            rows, aligns = block[1], block[2]
            cols = len(rows[0])
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for r_idx, row in enumerate(rows):
                for c_idx in range(cols):
                    cell = table.cell(r_idx, c_idx)
                    text = row[c_idx]
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    _docx_runs(par, text, size=10)
                    if r_idx == 0:
                        for run in par.runs:
                            run.bold = True
                        _docx_shade(cell._tc.get_or_add_tcPr(), "DBE4F0")
                    elif r_idx % 2 == 0:
                        _docx_shade(cell._tc.get_or_add_tcPr(), "F7F9FC")
                    align = aligns[c_idx]
                    if align == "right" or (r_idx > 0 and align == "left" and _NUMERIC_CELL.match(text or "x")):
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    elif align == "center":
                        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            header_props = table.rows[0]._tr.get_or_add_trPr()
            header_props.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        else:
            _docx_runs(doc.add_paragraph(), block[1])
    doc.save(str(path))


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
_PDF_FONTS = None


def _pdf_fonts():
    """Register DejaVu (ships with matplotlib) for full Unicode coverage;
    fall back to the built-in Helvetica family if it is somehow missing."""
    global _PDF_FONTS
    if _PDF_FONTS:
        return _PDF_FONTS
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        import matplotlib
        base = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        files = {"VSans": "DejaVuSans.ttf", "VSans-Bold": "DejaVuSans-Bold.ttf",
                 "VSans-Italic": "DejaVuSans-Oblique.ttf", "VSans-BoldItalic": "DejaVuSans-BoldOblique.ttf",
                 "VMono": "DejaVuSansMono.ttf"}
        for name, filename in files.items():
            pdfmetrics.registerFont(TTFont(name, str(base / filename)))
        pdfmetrics.registerFontFamily("VSans", normal="VSans", bold="VSans-Bold", italic="VSans-Italic", boldItalic="VSans-BoldItalic")
        _PDF_FONTS = {"regular": "VSans", "bold": "VSans-Bold", "mono": "VMono"}
    except Exception:
        _PDF_FONTS = {"regular": "Helvetica", "bold": "Helvetica-Bold", "mono": "Courier"}
    return _PDF_FONTS


def _pdf_markup(text, fonts):
    out = []
    for seg, flag in parse_inline(text):
        seg = xml_escape(seg)
        if flag == "b":
            out.append(f"<b>{seg}</b>")
        elif flag == "i":
            out.append(f"<i>{seg}</i>")
        elif flag == "c":
            out.append(f'<font name="{fonts["mono"]}" size="9" color="#B42318" backColor="#F3F4F6">{seg}</font>')
        else:
            out.append(seg)
    return "".join(out)


def write_pdf(path, content):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import HRFlowable, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

    blocks = parse_markdown_blocks(content)
    if not blocks:
        raise ValueError("Document has no content")
    fonts = _pdf_fonts()
    ink = colors.HexColor("#1F2937")
    base = ParagraphStyle("base", fontName=fonts["regular"], fontSize=10.5, leading=15, textColor=ink, spaceAfter=6)
    heading_sizes = {1: 21, 2: 15.5, 3: 12.5, 4: 11}
    heading_styles = {
        lvl: ParagraphStyle(f"h{lvl}", parent=base, fontName=fonts["bold"], fontSize=sz, leading=sz * 1.25,
                            spaceBefore=16 if lvl <= 2 else 10, spaceAfter=7, keepWithNext=1,
                            textColor=colors.HexColor("#0F172A"))
        for lvl, sz in heading_sizes.items()
    }
    code_style = ParagraphStyle("code", parent=base, fontName=fonts["mono"], fontSize=8.5, leading=11,
                                backColor=colors.HexColor("#F3F4F6"), borderPadding=6, spaceBefore=4, spaceAfter=12,
                                leftIndent=6, rightIndent=6)
    quote_style = ParagraphStyle("quote", parent=base, fontName=fonts["regular"], leftIndent=18, textColor=colors.HexColor("#4B5563"))
    cell_style = ParagraphStyle("cell", parent=base, fontSize=9, leading=12, spaceAfter=0)
    head_cell_style = ParagraphStyle("hcell", parent=cell_style, fontName=fonts["bold"], textColor=colors.white)

    margin = 0.9 * inch
    avail = letter[0] - 2 * margin
    title = _title_from_blocks(blocks)

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont(fonts["regular"], 8.5)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawCentredString(letter[0] / 2, 0.5 * inch, f"Page {doc_.page}")
        canvas.restoreState()

    story = []
    for block in blocks:
        tag = block[0]
        if tag == "heading":
            story.append(Paragraph(_pdf_markup(block[2], fonts), heading_styles[block[1]]))
        elif tag == "bullet":
            depth = block[2]
            style = ParagraphStyle("b", parent=base, leftIndent=18 + 16 * depth, bulletIndent=5 + 16 * depth, spaceAfter=3)
            story.append(Paragraph(_pdf_markup(block[1], fonts), style, bulletText="•" if depth == 0 else "–"))
        elif tag == "number":
            depth = block[3]
            style = ParagraphStyle("n", parent=base, leftIndent=22 + 16 * depth, bulletIndent=2 + 16 * depth, spaceAfter=3)
            story.append(Paragraph(_pdf_markup(block[1], fonts), style, bulletText=f"{block[2]}."))
        elif tag == "quote":
            story.append(Paragraph(f"<i>{_pdf_markup(block[1], fonts)}</i>", quote_style))
        elif tag == "hr":
            story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#CBD5E1"), spaceBefore=6, spaceAfter=10))
        elif tag == "code":
            wrapped = []
            for line in block[2].split("\n"):
                wrapped.extend(textwrap.wrap(line, 92, replace_whitespace=False, drop_whitespace=False,
                                             subsequent_indent="    ") or [""])
            story.append(Preformatted("\n".join(wrapped), code_style))
        elif tag == "table":
            rows, aligns = block[1], block[2]
            cols = len(rows[0])
            weights = [max(6, min(40, max(len(re.sub(r"[*`]", "", r[c])) for r in rows))) for c in range(cols)]
            total = sum(weights)
            widths = [max(avail * 0.07, avail * w / total) for w in weights]
            scale = avail / sum(widths)
            widths = [w * scale for w in widths]
            data = []
            for r_idx, row in enumerate(rows):
                data.append([Paragraph(_pdf_markup(cell, fonts), head_cell_style if r_idx == 0 else
                                       ParagraphStyle("c", parent=cell_style,
                                                      alignment=TA_RIGHT if aligns[c] == "right" or (aligns[c] == "left" and _NUMERIC_CELL.match(cell or "x"))
                                                      else TA_CENTER if aligns[c] == "center" else 0))
                             for c, cell in enumerate(row)])
            table = Table(data, colWidths=widths, repeatRows=1)
            style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                     ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                     ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                     ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]
            for r_idx in range(2, len(rows), 2):
                style.append(("BACKGROUND", (0, r_idx), (-1, r_idx), colors.HexColor("#F7F9FC")))
            table.setStyle(TableStyle(style))
            story.extend([table, Spacer(1, 12)])
        else:
            story.append(Paragraph(_pdf_markup(block[1], fonts), base))
    doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=margin, rightMargin=margin, topMargin=0.85 * inch,
                            bottomMargin=0.85 * inch, title=title, author="Velaris")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------
_UNSAFE_FORMULA = re.compile(
    r"\||\b(?:WEBSERVICE|HYPERLINK|CALL|REGISTER|EXEC|SHELL|IMPORTDATA|IMPORTXML|IMPORTHTML|IMPORTRANGE|FILTERXML|RTD|DDE)\s*\("
    r"|\[[^\]]*\]|https?://|file:|\\\\", re.IGNORECASE)
_PCT = re.compile(r"^[-+]?\d+(?:\.\d+)?%$")
_CUR = re.compile(r"^(\(|-)?\s*([$€£¥])\s*(\d{1,3}(?:,\d{3})*|\d+)(\.\d+)?\)?$")
_THOUSANDS = re.compile(r"^[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?$")
_PLAIN_NUM = re.compile(r"^[-+]?\d+(?:\.\d+)?$")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _coerce_cell(value):
    """-> (value, number_format|None, is_text_forced)"""
    if value is None or isinstance(value, (bool, int, float)):
        return value, None, False
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False), None, True
    s = strip_control(value).strip()
    if s == "":
        return None, None, False
    if s.startswith("="):
        depth_ok = s.count("(") == s.count(")")
        if depth_ok and not _UNSAFE_FORMULA.search(s) and len(s) < 1000:
            return s, None, False
        return s, None, True
    if _PCT.match(s):
        number = float(s[:-1]) / 100
        return number, "0.0%" if "." in s else "0%", False
    m = _CUR.match(s)
    if m:
        number = float(m.group(3).replace(",", "") + (m.group(4) or ""))
        if m.group(1):
            number = -number
        return number, f'"{m.group(2)}"#,##0.00', False
    if _THOUSANDS.match(s):
        number = float(s.replace(",", ""))
        return (int(number) if number.is_integer() and "." not in s else number), "#,##0.##" if "." in s else "#,##0", False
    if _PLAIN_NUM.match(s):
        digits = s.lstrip("+-")
        if (len(digits) > 1 and digits.startswith("0") and not digits.startswith("0.")) or len(digits) > 15:
            return s, None, True  # ZIP codes, IDs, long numbers stay text
        return (float(s) if "." in s else int(s)), None, False
    if _ISO_DATE.match(s):
        try:
            return dt.date.fromisoformat(s), "yyyy-mm-dd", False
        except ValueError:
            pass
    return s[:32000], None, False


def _sheet_title(name, taken):
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", str(name)).strip()[:31] or "Sheet"
    candidate, n = cleaned, 2
    while candidate.lower() in {t.lower() for t in taken}:
        suffix = f" {n}"
        candidate = cleaned[:31 - len(suffix)] + suffix; n += 1
    return candidate


def _rows_from(value):
    """Normalize one sheet's data (list of lists / list of dicts / scalars) to a list of lists."""
    if isinstance(value, dict) and "rows" in value:
        value = value["rows"]
    if not isinstance(value, list):
        raise ValueError("Sheet data must be a list of rows")
    if value and all(isinstance(r, dict) for r in value):
        keys = []
        for r in value:
            for k in r:
                if k not in keys:
                    keys.append(k)
        return [keys] + [[r.get(k) for k in keys] for r in value]
    return [r if isinstance(r, (list, tuple)) else [r] for r in value]


def _parse_sheets(content):
    if isinstance(content, str):
        try:
            data = json.loads(content)
        except ValueError:
            data = list(csv.reader(io.StringIO(content)))  # model sent CSV instead of JSON
    else:
        data = content
    sheets = []
    if isinstance(data, dict):
        if isinstance(data.get("sheets"), list):
            for i, s in enumerate(data["sheets"]):
                sheets.append((s.get("name", f"Sheet{i + 1}") if isinstance(s, dict) else f"Sheet{i + 1}", _rows_from(s)))
        elif isinstance(data.get("sheets"), dict):
            sheets = [(k, _rows_from(v)) for k, v in data["sheets"].items()]
        elif "rows" in data:
            sheets = [(data.get("name", "Sheet1"), _rows_from(data))]
        elif data and all(isinstance(v, list) for v in data.values()):
            sheets = [(k, _rows_from(v)) for k, v in data.items()]
        else:
            raise ValueError("Unrecognized spreadsheet structure")
    elif isinstance(data, list):
        if data and all(isinstance(r, dict) and "rows" in r for r in data):
            sheets = [(r.get("name", f"Sheet{i + 1}"), _rows_from(r)) for i, r in enumerate(data)]
        else:
            sheets = [("Sheet1", _rows_from(data))]
    else:
        raise ValueError("Spreadsheet content must be JSON rows")
    sheets = [(n, r) for n, r in sheets if r]
    if not sheets:
        raise ValueError("Spreadsheet has no data")
    return sheets[:MAX_SHEETS]


def write_xlsx(path, content):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    sheets = _parse_sheets(content)
    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    zebra = PatternFill("solid", fgColor="F7F9FC")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    taken = []
    for name, rows in sheets:
        title = _sheet_title(name, taken); taken.append(title)
        ws = wb.create_sheet(title)
        width = max(len(r) for r in rows[:MAX_SHEET_ROWS])
        widths = [8] * width
        for r_idx, row in enumerate(rows[:MAX_SHEET_ROWS], start=1):
            for c_idx in range(1, width + 1):
                raw = row[c_idx - 1] if c_idx - 1 < len(row) else None
                value, fmt, force_text = _coerce_cell(raw)
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = value
                if force_text and isinstance(value, str):
                    cell.data_type = "s"
                if fmt:
                    cell.number_format = fmt
                cell.border = border
                shown = str(value) if value is not None else ""
                if isinstance(value, str) and value.startswith("=") and not force_text:
                    shown = "0" * 10
                widths[c_idx - 1] = max(widths[c_idx - 1], min(len(shown) + 2, 50))
                if r_idx == 1:
                    cell.font = header_font; cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    cell.alignment = Alignment(vertical="top", wrap_text=isinstance(value, str) and len(value) > 50)
                    if r_idx % 2 == 1:
                        cell.fill = zebra
        for c_idx, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c_idx)].width = max(w, 10)
        ws.freeze_panes = "A2"
        if len(rows) > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(width)}{min(len(rows), MAX_SHEET_ROWS)}"
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    wb.properties.creator = "Velaris"
    wb.save(str(path))


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------
_PPT_NAVY, _PPT_ACCENT, _PPT_TEXT, _PPT_MUTED = "0F172A", "4F46E5", "1F2937", "64748B"


def _rgb(hex_):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hex_)


def _slide_lines(body):
    """body (str | list) -> [(text, depth)]"""
    if isinstance(body, list):
        raw = []
        for item in body:
            raw.extend(str(item).split("\n"))
    else:
        raw = str(body or "").replace("\\n", "\n").split("\n")
    items = []
    for line in raw:
        expanded = line.replace("\t", "    ")
        text = expanded.strip()
        if not text or re.match(r"^#{1,6}\s*$", text):
            continue
        depth = min((len(expanded) - len(expanded.lstrip())) // 2, 2)
        text = re.sub(r"^(?:[-*+•–]\s+|\d+[.)]\s+|#{1,6}\s+)", "", text)
        items.append((text, depth))
    return items


def _est_lines(text, size_pt, width_in):
    chars_per_line = max(8, int(width_in * 72 / (size_pt * 0.5)))
    return max(1, math.ceil(len(text) / chars_per_line))


def _fits(items, size_pt, width_in=11.4, height_in=4.9):
    total = 0.0
    for text, depth in items:
        lines = _est_lines(text, size_pt, width_in - 0.45 - 0.4 * depth)
        total += lines * size_pt * 1.22 / 72 + 0.13
    return total <= height_in


def _pick_size(items):
    for size in (26, 24, 22, 20, 18, 16, 14):
        if _fits(items, size):
            return size
    return 14


def _chunk_items(items, min_size=18):
    chunks, current = [], []
    for item in items:
        if current and not _fits(current + [item], min_size):
            chunks.append(current); current = []
        current.append(item)
    if current:
        chunks.append(current)
    return chunks or [[]]


def _set_bullet(paragraph, depth):
    from pptx.oxml.ns import qn
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(342900 + depth * 342900)))
    pPr.set("indent", str(-285750))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bullet = pPr.makeelement(qn("a:buChar"), {"char": "•" if depth == 0 else "–"})
    pPr.append(bullet)


def _ppt_runs(paragraph, text, size, color):
    from pptx.util import Pt
    for seg, flag in parse_inline(text):
        if not seg:
            continue
        run = paragraph.add_run()
        run.text = seg
        run.font.size = Pt(size)
        run.font.color.rgb = _rgb(color)
        run.font.bold = True if flag == "b" else None
        run.font.italic = True if flag == "i" else None


def _style_title(shape, text, size, color, left, top, width, height, bold=True, align=None, anchor=None):
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt
    shape.left, shape.top, shape.width, shape.height = Inches(left), Inches(top), Inches(width), Inches(height)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = anchor or MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = _rgb(color)


def _background(slide, hex_):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = _rgb(hex_)


def _accent_bar(slide, left, top, width, height, color=_PPT_ACCENT):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    bar.fill.solid(); bar.fill.fore_color.rgb = _rgb(color)
    bar.line.fill.background()
    return bar


def _slide_number(slide, number):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(11.9), Inches(6.9), Inches(1.0), Inches(0.4))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run(); run.text = str(number)
    run.font.size = Pt(11); run.font.color.rgb = _rgb(_PPT_MUTED)


def write_pptx(path, content):
    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    data = content
    if isinstance(data, str):
        data = json.loads(data)
    if isinstance(data, dict):
        data = data.get("slides") or [data]
    if not isinstance(data, list) or not data:
        raise ValueError("Presentation needs a non-empty list of slides")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout_title, layout_title_only = prs.slide_layouts[0], prs.slide_layouts[5]
    number = 0

    def add_title_slide(title, subtitle, notes):
        nonlocal number
        slide = prs.slides.add_slide(layout_title)
        number += 1
        _background(slide, _PPT_NAVY)
        _accent_bar(slide, 0.9, 2.15, 1.3, 0.09)
        _style_title(slide.shapes.title, title, 44, "FFFFFF", 0.9, 2.35, 11.5, 1.7, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
        sub = slide.placeholders[1]
        if subtitle:
            _style_title(sub, subtitle, 22, "CBD5E1", 0.9, 4.15, 11.5, 1.3, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
        else:
            sub._element.getparent().remove(sub._element)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    def add_section_slide(title, subtitle, notes):
        nonlocal number
        slide = prs.slides.add_slide(layout_title_only)
        number += 1
        _background(slide, _PPT_ACCENT)
        _style_title(slide.shapes.title, title, 40, "FFFFFF", 0.9, 2.6, 11.5, 1.6, align=PP_ALIGN.LEFT)
        if subtitle:
            box = slide.shapes.add_textbox(Inches(0.9), Inches(4.2), Inches(11.5), Inches(1.0))
            box.text_frame.word_wrap = True
            _ppt_runs(box.text_frame.paragraphs[0], subtitle, 20, "E0E7FF")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    def new_content_slide(title, notes):
        nonlocal number
        slide = prs.slides.add_slide(layout_title_only)
        number += 1
        _background(slide, "FFFFFF")
        _accent_bar(slide, 0, 0.5, 0.22, 0.95)
        _style_title(slide.shapes.title, title, 32 if len(title) < 55 else 26, _PPT_NAVY, 0.7, 0.45, 11.9, 1.05,
                     align=PP_ALIGN.LEFT)
        _slide_number(slide, number)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    for index, raw in enumerate(data):
        slide_data = raw if isinstance(raw, dict) else {"title": str(raw)}
        title = strip_control(slide_data.get("title") or f"Slide {index + 1}")
        subtitle = strip_control(slide_data.get("subtitle") or "")
        notes = strip_control(slide_data.get("notes") or slide_data.get("speaker_notes") or "")
        items = _slide_lines(slide_data.get("bullets") if slide_data.get("bullets") is not None else slide_data.get("body"))
        table = slide_data.get("table")
        layout = str(slide_data.get("layout", "")).lower()
        if layout == "title" or (index == 0 and not items and not table and layout != "section"):
            add_title_slide(title, subtitle, notes); continue
        if layout == "section" or (not items and not table):
            add_section_slide(title, subtitle, notes); continue

        if table:
            rows = _rows_from(table)
            cols = max(len(r) for r in rows)
            header, body_rows = rows[0], rows[1:]
            per_slide = 8
            groups = [body_rows[i:i + per_slide] for i in range(0, max(len(body_rows), 1), per_slide)] or [[]]
            for g_idx, group in enumerate(groups):
                slide = new_content_slide(title if g_idx == 0 else f"{title} (cont.)", notes if g_idx == 0 else "")
                data_rows = [header] + group
                shape = slide.shapes.add_table(len(data_rows), cols, Inches(0.9), Inches(1.75), Inches(11.5),
                                               Inches(0.5 * len(data_rows)))
                for r_idx, row in enumerate(data_rows):
                    for c_idx in range(cols):
                        cell = shape.table.cell(r_idx, c_idx)
                        cell.text = strip_control(str(row[c_idx]) if c_idx < len(row) and row[c_idx] is not None else "")
                        for par in cell.text_frame.paragraphs:
                            for run in par.runs:
                                run.font.size = Pt(15 if cols <= 4 else 12)
                                run.font.bold = True if r_idx == 0 else None
            continue

        for c_idx, chunk in enumerate(_chunk_items(items)):
            slide = new_content_slide(title if c_idx == 0 else f"{title} (cont.)", notes if c_idx == 0 else "")
            box = slide.shapes.add_textbox(Inches(0.9), Inches(1.75), Inches(11.4), Inches(4.95))
            frame = box.text_frame
            frame.word_wrap = True
            size = _pick_size(chunk)
            for p_idx, (text, depth) in enumerate(chunk):
                paragraph = frame.paragraphs[0] if p_idx == 0 else frame.add_paragraph()
                paragraph.space_after = Pt(8)
                _ppt_runs(paragraph, strip_control(text), size - 2 * depth, _PPT_TEXT)
                _set_bullet(paragraph, depth)
    prs.core_properties.title = strip_control(str((data[0] or {}).get("title", "Presentation")) if isinstance(data[0], dict) else "Presentation")[:120]
    prs.core_properties.author = "Velaris"
    prs.save(str(path))


# --------------------------------------------------------------------------
# Charts
# --------------------------------------------------------------------------
CHART_COLORS = ["#4F46E5", "#10B981", "#F59E0B", "#EF4444", "#3B82F6", "#8B5CF6", "#14B8A6", "#EC4899"]


def _num(v):
    try:
        f = float(str(v).replace(",", "").replace("%", "").replace("$", "")) if not isinstance(v, (int, float)) else float(v)
        return f if math.isfinite(f) else float("nan")
    except (TypeError, ValueError):
        return float("nan")


def render_chart(spec, path):
    """Real matplotlib chart on a white background (readable anywhere).
    Uses the Figure object API: pyplot's global state is not thread-safe."""
    from matplotlib.figure import Figure
    if not isinstance(spec, dict):
        raise ValueError("Chart spec must be a JSON object")
    chart_type = str(spec.get("type", "bar")).lower()
    labels = [str(x) for x in (spec.get("labels") or [])]
    series = spec.get("series")
    if not series:
        series = [{"name": "Series 1", "values": spec.get("values", [])}]
    series = [s for s in series if isinstance(s, dict)]
    if not series or not any(s.get("values") or s.get("points") for s in series):
        raise ValueError("Chart spec has no data")

    fig = Figure(figsize=(8, 4.8), dpi=160, facecolor="white")
    ax = fig.add_subplot(111)
    ax.set_facecolor("white")

    def vals(s):
        return [_num(v) for v in (s.get("values") or [])]

    if chart_type == "pie":
        values = [max(v, 0) if not math.isnan(v) else 0 for v in vals(series[0])]
        if not sum(values):
            raise ValueError("Pie chart needs positive values")
        ax.pie(values, labels=labels[:len(values)] or None, autopct="%1.0f%%", colors=CHART_COLORS * 3,
               startangle=90, wedgeprops={"edgecolor": "white"}, textprops={"color": "#1F2937"})
        ax.axis("equal")
    elif chart_type == "scatter":
        for i, s in enumerate(series):
            if s.get("points"):
                xs, ys = zip(*[(_num(p[0]), _num(p[1])) for p in s["points"]])
            else:
                ys = vals(s)
                xs = [_num(x) for x in labels] if labels and all(not math.isnan(_num(x)) for x in labels) else list(range(len(ys)))
            ax.scatter(xs, ys, label=s.get("name", f"Series {i + 1}"), color=CHART_COLORS[i % len(CHART_COLORS)], s=36, alpha=0.85)
        ax.grid(alpha=0.25)
    elif chart_type in ("line", "area"):
        count = max(len(vals(s)) for s in series)
        xs = labels[:count] if len(labels) >= count else [str(j + 1) for j in range(count)]
        for i, s in enumerate(series):
            ys = vals(s)
            color = CHART_COLORS[i % len(CHART_COLORS)]
            ax.plot(xs[:len(ys)], ys, marker="o", linewidth=2, label=s.get("name", f"Series {i + 1}"), color=color)
            if chart_type == "area":
                ax.fill_between(xs[:len(ys)], ys, alpha=0.15, color=color)
        ax.grid(alpha=0.25)
        if count > 8:
            ax.tick_params(axis="x", rotation=30)
    else:  # bar / hbar, grouped or stacked
        count = len(labels) if labels else max(len(vals(s)) for s in series)
        xs_labels = labels[:count] if labels else [str(j + 1) for j in range(count)]
        horizontal = chart_type in ("hbar", "barh")
        stacked = bool(spec.get("stacked"))
        width = 0.8 if stacked else 0.8 / max(1, len(series))
        bottoms = [0.0] * count
        for i, s in enumerate(series):
            ys = (vals(s) + [0.0] * count)[:count]
            ys = [0.0 if math.isnan(y) else y for y in ys]
            pos = [j + (0 if stacked else i * width) for j in range(count)]
            color = CHART_COLORS[i % len(CHART_COLORS)]
            name = s.get("name", f"Series {i + 1}")
            if horizontal:
                bars = ax.barh(pos, ys, height=width, left=bottoms if stacked else None, label=name, color=color)
            else:
                bars = ax.bar(pos, ys, width=width, bottom=bottoms if stacked else None, label=name, color=color)
            if len(series) == 1 and count <= 14:
                ax.bar_label(bars, fmt="%g", padding=3, fontsize=8, color="#374151")
            if stacked:
                bottoms = [b + y for b, y in zip(bottoms, ys)]
        offset = 0 if stacked else (len(series) - 1) * width / 2
        ticks = [j + offset for j in range(count)]
        if horizontal:
            ax.set_yticks(ticks); ax.set_yticklabels(xs_labels); ax.grid(axis="x", alpha=0.25)
        else:
            ax.set_xticks(ticks); ax.set_xticklabels(xs_labels, rotation=20 if count > 4 else 0, ha="right" if count > 4 else "center")
            ax.grid(axis="y", alpha=0.25)
        ax.set_axisbelow(True)
    if len(series) > 1 and chart_type != "pie":
        ax.legend(frameon=False)
    if chart_type != "pie":
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
    if spec.get("title"):
        ax.set_title(str(spec["title"]), fontsize=13, fontweight="bold", color="#0F172A", pad=12)
    if spec.get("x_label") and chart_type != "pie":
        ax.set_xlabel(str(spec["x_label"]))
    if spec.get("y_label") and chart_type != "pie":
        ax.set_ylabel(str(spec["y_label"]))
    fig.tight_layout()
    fig.savefig(str(path), facecolor="white")


# --------------------------------------------------------------------------
# Images, SVG, base64, text helpers
# --------------------------------------------------------------------------
_ASPECTS = {"portrait": (832, 1216), "tall": (832, 1216), "9:16": (832, 1216), "3:4": (832, 1216),
            "landscape": (1216, 832), "wide": (1216, 832), "16:9": (1216, 832), "4:3": (1216, 832)}


def write_image(path, content):
    from PIL import Image
    prompt, aspect = str(content), "square"
    parsed = content if isinstance(content, dict) else None
    if parsed is None:
        try:
            candidate = json.loads(content)
            parsed = candidate if isinstance(candidate, dict) else None
        except (TypeError, ValueError):
            parsed = None
    if parsed and "prompt" in parsed:
        prompt, aspect = str(parsed["prompt"]), str(parsed.get("aspect", "square")).lower()
    width, height = _ASPECTS.get(aspect, (1024, 1024))
    prompt = strip_control(prompt).strip()[:800] or "abstract art"
    url = IMAGE_API_URL.format(prompt=quote(prompt)) + f"?width={width}&height={height}&nologo=true&model=flux&seed={uuid.uuid4().int % 10**6}"
    last_error = None
    for attempt in range(2):
        try:
            response = requests.get(url, timeout=IMAGE_TIMEOUT, headers={"User-Agent": "Velaris/2"})
            response.raise_for_status()
            image = Image.open(io.BytesIO(response.content))
            image.load()
            break
        except Exception as error:  # network error, HTTP error, or bytes that are not an image
            last_error = error
    else:
        raise ValueError(f"Image generation failed ({type(last_error).__name__}: {last_error}). Try again in a moment.")
    ext = path.suffix.lower()
    if ext in (".jpg", ".jpeg"):
        image.convert("RGB").save(str(path), "JPEG", quality=92)
    elif ext == ".webp":
        image.save(str(path), "WEBP", quality=92)
    elif ext == ".gif":
        image.convert("P").save(str(path), "GIF")
    else:
        image.save(str(path), "PNG")


def _clean_svg(text):
    if "<svg" not in text.lower():
        raise ValueError("SVG content does not contain an <svg> element")
    if re.search(r"<!ENTITY|<!DOCTYPE[^>]*\[", text, re.IGNORECASE):
        raise ValueError("SVG with custom entities is not allowed")
    text = re.sub(r"<script\b.*?</script\s*>", "", text, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r"<foreignObject\b.*?</foreignObject\s*>", "", text, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r"\son\w+\s*=\s*(\"[^\"]*\"|'[^']*')", "", text, flags=re.IGNORECASE)
    text = re.sub(r"(href\s*=\s*[\"'])\s*javascript:[^\"']*", r"\1#", text, flags=re.IGNORECASE)
    text = text.strip()
    if text.startswith("<svg") and "xmlns=" not in text.split(">", 1)[0]:
        text = text.replace("<svg", '<svg xmlns="http://www.w3.org/2000/svg"', 1)
    try:
        ET.fromstring(text.encode("utf-8"))
    except ET.ParseError as error:
        raise ValueError(f"SVG is not well-formed XML: {error}")
    return text


_FENCE_WRAP = re.compile(r"^\s*```[\w+.\-]*[ \t]*\n(.*?)\n[ \t]*```\s*$", re.DOTALL)


def _write_text(path, content):
    warnings = []
    if isinstance(content, (dict, list)):
        content = json.dumps(content, indent=2, ensure_ascii=False)
    text = strip_control(content).replace("\r\n", "\n")
    ext = path.suffix.lower()
    if ext not in (".md", ".markdown"):
        wrapped = _FENCE_WRAP.match(text)
        if wrapped and "```" not in wrapped.group(1):
            text = wrapped.group(1)
    if "\n" not in text and text.count("\\n") >= 2 and len(text) > 60:
        text = text.replace("\\n", "\n").replace("\\t", "\t").replace('\\"', '"')
    if ext == ".json":
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except ValueError as error:
            warnings.append(f"⚠️ '{path.name}' is not valid JSON ({error}); saved as written.")
    elif ext == ".svg":
        text = _clean_svg(text)
    if not text.endswith("\n"):
        text += "\n"
    path.write_text(text, encoding="utf-8")
    return warnings


def _write_base64(path, content):
    data = re.sub(r"^data:[^,]*,", "", str(content).strip())
    data = re.sub(r"\s+", "", data)
    data += "=" * (-len(data) % 4)
    if len(data) * 3 // 4 > MAX_BASE64_BYTES:
        raise ValueError("Binary payload is too large")
    path.write_bytes(base64.b64decode(data, validate=True))


def _parse_json_content(content):
    if isinstance(content, (dict, list)):
        return content
    text = strip_control(content).strip()
    wrapped = _FENCE_WRAP.match(text)
    if wrapped:
        text = wrapped.group(1)
    return json.loads(text)


# --------------------------------------------------------------------------
# Writer dispatch, validation, public entry point
# --------------------------------------------------------------------------
def _write_zip(root, path, content, prompt, depth):
    if depth >= 1:
        raise ValueError("Nested ZIP archives are not supported")
    entries = _parse_json_content(content)
    if isinstance(entries, dict):
        entries = entries.get("files", [])
    if not isinstance(entries, list):
        raise ValueError("ZIP content must be a JSON list of files")
    warnings, written = [], 0
    build_dir = Path(tempfile.mkdtemp(prefix=".zip-build-", dir=root))
    try:
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for entry in entries[:MAX_ZIP_ENTRIES]:
                try:
                    item = normalize_artifact_item(entry, "")
                    if item["kind"] == "zip":
                        raise ValueError("nested ZIP skipped")
                    item["path"] = _unique_rel(build_dir, item["path"])
                    warnings += _write_item(build_dir, item, prompt, depth + 1)
                    validate_artifact(build_dir / item["path"], item["kind"])
                    archive.write(build_dir / item["path"], item["path"])
                    written += 1
                except Exception as error:
                    name = entry.get("path", "?") if isinstance(entry, dict) else "?"
                    warnings.append(f"⚠️ Skipped '{name}' inside the ZIP: {error}")
    finally:
        shutil.rmtree(build_dir, ignore_errors=True)
    if not written:
        raise ValueError("ZIP archive ended up with no valid files")
    return warnings


def _write_item(root, item, prompt="", depth=0):
    path = root / safe_path(item["path"])
    path.parent.mkdir(parents=True, exist_ok=True)
    kind, content = item["kind"], item.get("content", "")
    if kind == "text":
        return _write_text(path, content)
    if kind == "base64":
        _write_base64(path, content)
    elif kind == "image":
        if path.suffix.lower() == ".svg":
            return _write_text(path, content)
        write_image(path, content)
    elif kind == "chart":
        render_chart(_parse_json_content(content), path)
    elif kind == "docx":
        write_docx(path, content)
    elif kind == "xlsx":
        write_xlsx(path, content)
    elif kind == "pptx":
        write_pptx(path, content)
    elif kind == "pdf":
        write_pdf(path, content)
    elif kind == "zip":
        return _write_zip(root, path, content, prompt, depth)
    elif kind == "stl":
        triangles, warnings = run_stl_program(_parse_json_content(content))
        if not triangles:
            raise ValueError("The 3D build program produced no geometry")
        path.write_text(render_ascii_stl(triangles), encoding="ascii")
        return warnings
    else:
        raise ValueError(f"Unsupported artifact kind: {kind}")
    return []


def validate_artifact(path, kind):
    """Re-open the produced file with the matching library to prove it is real."""
    path = Path(path)
    data = path.read_bytes()
    if not data:
        raise ValueError("Generated file is empty")
    ext = path.suffix.lower()
    if kind == "pdf":
        if not data.startswith(b"%PDF-") or b"%%EOF" not in data[-2048:]:
            raise ValueError("Generated PDF is incomplete")
    elif kind in ("docx", "xlsx", "pptx"):
        if not zipfile.is_zipfile(path):
            raise ValueError(f"Generated {kind.upper()} is not a valid Office container")
        if kind == "docx":
            from docx import Document
            doc = Document(str(path))
            if not doc.paragraphs and not doc.tables:
                raise ValueError("Generated DOCX is empty")
        elif kind == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(path))
            if not any(ws.max_row >= 1 and ws.max_column >= 1 for ws in wb.worksheets):
                raise ValueError("Generated XLSX has no cells")
        else:
            from pptx import Presentation
            if len(Presentation(str(path)).slides) < 1:
                raise ValueError("Generated PPTX has no slides")
    elif kind in ("image", "chart") and ext in RASTER_EXTENSIONS:
        from PIL import Image
        with Image.open(path) as image:
            image.verify()
    elif kind == "zip":
        with zipfile.ZipFile(path) as archive:
            if archive.testzip() is not None or not archive.namelist():
                raise ValueError("Generated ZIP is corrupt or empty")
    elif kind == "stl":
        head = data[:2000].lower()
        facets = data.lower().count(b"facet normal")
        if not head.startswith(b"solid") or facets < 4 or b"nan" in data.lower():
            raise ValueError("Generated STL is empty or malformed")
    return len(data)


def produce_artifact(root, raw_item, prompt=""):
    """Normalize -> write -> validate one artifact. Returns
    {'path', 'kind', 'bytes', 'warnings'}; raises ValueError on failure and
    never leaves a partial file behind."""
    root = Path(root)
    item = normalize_artifact_item(raw_item, prompt)
    item["path"] = _unique_rel(root, item["path"])
    target = root / item["path"]
    try:
        warnings = _write_item(root, item, prompt)
        size = validate_artifact(target, item["kind"])
    except json.JSONDecodeError as error:
        target.unlink(missing_ok=True)
        raise ValueError(f"its content was not valid JSON ({error.msg} at position {error.pos})")
    except Exception:
        target.unlink(missing_ok=True)
        raise
    return {"path": item["path"], "kind": item["kind"], "bytes": size, "warnings": warnings}


# --------------------------------------------------------------------------
# Last-resort fallback: build the requested file from the chat reply itself
# --------------------------------------------------------------------------
def _slug(text, default="velaris-output"):
    slug = re.sub(r"[^a-z0-9]+", "-", str(text).lower()).strip("-")[:48].strip("-")
    return slug or default


def fallback_item_from_reply(kind, reply, prompt):
    """Return a file item built from the assistant's reply text, or None when
    the reply is too thin / the kind cannot be derived from prose."""
    reply = str(reply or "").strip()
    if len(reply) < 80:
        return None
    blocks = parse_markdown_blocks(reply)
    name = _slug(_title_from_blocks(blocks, "") or prompt)
    if kind in ("docx", "pdf"):
        return {"path": f"{name}{default_extension(kind)}", "kind": kind, "content": reply}
    if kind == "xlsx":
        for block in blocks:
            if block[0] == "table":
                rows = [[re.sub(r"[*`]", "", c) for c in row] for row in block[1]]
                return {"path": f"{name}.xlsx", "kind": "xlsx", "content": rows}
        return None
    if kind == "pptx":
        slides, current = [], None
        for block in blocks:
            if block[0] == "heading" and block[1] <= 2:
                current = {"title": re.sub(r"[*`]", "", block[2]), "body": []}
                slides.append(current)
            elif current is not None and block[0] in ("bullet", "number", "para"):
                current["body"].append(block[1])
        slides = [s for s in slides if s["body"]]
        if not slides:
            return None
        first = {"title": slides[0]["title"], "subtitle": "Generated by Velaris", "layout": "title"}
        return {"path": f"{name}.pptx", "kind": "pptx", "content": [first] + [{**s, "body": "\n".join(s["body"])} for s in slides]}
    return None
